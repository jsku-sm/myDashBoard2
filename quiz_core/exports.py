"""Teacher-authorized quiz handouts. No student records, embedded fonts or network."""
from __future__ import annotations
from copy import deepcopy
from html import escape
from io import BytesIO
import re
import zipfile

DOCX_MIME="application/vnd.openxmlformats-officedocument.wordprocessingml.document"

def filename(value, limit=70):
    result=re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",str(value))
    result=re.sub(r"\s+"," ",result).strip(" .")
    return (result[:limit] or "퀴즈")

def _check(papers, include_answers):
    if not papers or not any(p.get("questions") for p in papers):
        raise ValueError("내려받을 문항이 없습니다.")
    if type(include_answers) is not bool:
        raise ValueError("정답 포함 여부를 확인하세요.")
    for paper in papers:
        for q in paper["questions"]:
            if not isinstance(q.get("prompt"),str) or not isinstance(q.get("options"),list) or len(q["options"])!=4:
                raise ValueError("문제 또는 보기 형식이 올바르지 않습니다.")
            if include_answers and (type(q.get("answer")) is not int or not 0 <= q["answer"] < 4):
                raise ValueError("정답을 포함할 수 없는 문항이 있습니다.")

def _docx(papers, include_answers):
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ValueError("Word 파일 생성에는 python-docx가 필요합니다. 수정본 requirements.txt를 함께 적용해 주세요.") from exc
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.6)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    section.header_distance=Cm(.6);section.footer_distance=Cm(.7)
    for style_name in ("Normal","Title","Heading 1","Heading 2"):
        style=doc.styles[style_name]
        style.font.name="맑은 고딕"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"맑은 고딕")
        style.font.size=Pt(11)
        style.paragraph_format.line_spacing=1.35
        style.paragraph_format.space_after=Pt(5)
    doc.styles["Title"].font.size=Pt(20); doc.styles["Title"].font.bold=True
    doc.styles["Title"].font.color.rgb=RGBColor.from_string("263750")
    doc.styles["Heading 1"].font.size=Pt(15)
    doc.styles["Heading 2"].font.size=Pt(12)
    doc.core_properties.author="수업 대시보드"
    doc.core_properties.title="개념확인 퀴즈"
    doc.core_properties.subject="교사가 출제한 문제 모음"
    header=section.header.paragraphs[0]
    header.text="CONCEPT CHECK   |   수업 대시보드"
    header.runs[0].font.size=Pt(9)
    header.runs[0].font.color.rgb=RGBColor.from_string("60728F")
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    footer.add_run("교사용 · 정답 포함   " if include_answers else "학생용 · 문제만   ")
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    for paper_index,paper in enumerate(papers):
        if paper_index: doc.add_page_break()
        doc.add_paragraph(paper["title"],"Title")
        meta=doc.add_paragraph(f'{paper.get("subject","")}  |  {paper.get("unit","")}  |  {len(paper["questions"])}문항')
        meta.runs[0].font.color.rgb=RGBColor.from_string("63748F")
        mode="교사용 · 정답과 해설은 뒤쪽에 있습니다." if include_answers else "학생용 · 각 문항에서 알맞은 보기를 하나 고르세요."
        doc.add_paragraph(mode)
        doc.add_paragraph("학급: __________   이름: _______________   날짜: _______________")
        for q in paper["questions"]:
            head=doc.add_paragraph(f'문항 {q.get("number",1):02d}',"Heading 2")
            head.paragraph_format.space_before=Pt(14)
            head.paragraph_format.keep_with_next=True
            p=doc.add_paragraph(q["prompt"])
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
            table=doc.add_table(rows=4, cols=2)
            table.autofit=False
            table.columns[0].width=Cm(.9);table.columns[1].width=Cm(16)
            table.style="Table Grid"
            for j,opt in enumerate(q["options"]):
                table.cell(j,0).text=str(j+1)
                table.cell(j,1).text=opt
                cell=table.cell(j,0)
                shading=OxmlElement("w:shd");shading.set(qn("w:fill"),"EEF2F9")
                cell._tc.get_or_add_tcPr().append(shading)
                for c in table.rows[j].cells:
                    for paragraph in c.paragraphs:
                        paragraph.paragraph_format.space_after=Pt(4)
                        paragraph.paragraph_format.space_before=Pt(4)
                        paragraph.paragraph_format.keep_with_next=(j<3)
                trpr=table.rows[j]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
            note=doc.add_paragraph("나의 생각: __________________________________________________________")
            note.paragraph_format.space_before=Pt(5)
            note.paragraph_format.space_after=Pt(9)
            note.runs[0].font.size=Pt(9)
            note.runs[0].font.color.rgb=RGBColor.from_string("66788F")
        if include_answers:
            doc.add_page_break()
            doc.add_paragraph("정답 및 해설","Title")
            doc.add_paragraph(paper["title"])
            doc.add_paragraph("교사용 · 학생에게 배포하기 전에 정답 포함 여부를 확인하세요.")
            for q in paper["questions"]:
                doc.add_paragraph(f'{q.get("number",1):02d}번  |  정답 {q["answer"]+1}번',"Heading 2")
                doc.add_paragraph(q["options"][q["answer"]])
                if q.get("explanation"):doc.add_paragraph(q["explanation"])
    out=BytesIO();doc.save(out)
    return out.getvalue()

def _html(papers, include_answers):
    e=lambda value:escape(str(value),quote=True)
    css="""*{box-sizing:border-box}body{font-family:'Apple SD Gothic Neo','Malgun Gothic',sans-serif;background:#f4f6fb;color:#273650;line-height:1.75;margin:0}
main{max-width:900px;margin:28px auto;background:white;padding:36px;border-radius:18px}.toolbar{max-width:900px;margin:18px auto;font-size:14px}
button{padding:10px 18px;background:#485f91;color:white;border:0;border-radius:9px;cursor:pointer}h1{font-size:25px;line-height:1.5}
.meta{color:#60728d}.question{margin:26px 0;padding:22px;border:1px solid #dce3ef;border-radius:15px;break-inside:avoid}
.number{color:#766197;font-weight:800}.prompt{font-size:18px;white-space:pre-wrap;background:#f5f7fc;padding:13px;border-radius:10px}
.options{list-style:none;padding:0;display:grid;grid-template-columns:1fr 1fr;gap:10px}.options li{border:1px solid #dce3ef;padding:12px;border-radius:10px;white-space:pre-wrap;overflow-wrap:anywhere}
.answer-section{border-top:2px solid #bec9dc;margin-top:30px;padding-top:15px}.answer{break-inside:avoid}
.paper+.paper{margin-top:50px;break-before:page}.workspace{font-size:13px;color:#708099;border-bottom:1px solid #dce3ef;padding:12px 0}
@media(max-width:600px){main{padding:18px;margin:8px}.options{grid-template-columns:1fr}.question{padding:14px}}
@media print{@page{size:A4;margin:17mm}body{background:white}.toolbar{display:none}main{margin:0;max-width:none;padding:0;border:0}.paper+.paper{break-before:page}.question{box-shadow:none;padding:14px}.answer-section{break-before:page}}
"""
    out=['<!doctype html><html lang="ko"><head><meta charset="utf-8">',
         '<meta name="viewport" content="width=device-width,initial-scale=1">',
         '<title>개념확인 퀴즈</title><style>',css,'</style></head><body>',
         '<div class="toolbar"><button id="print-quiz">인쇄 / PDF로 저장</button>  브라우저 인쇄 메뉴에서 PDF 저장을 선택할 수 있습니다.</div><main>']
    for paper in papers:
        out.append(f'<section class="paper"><p class="meta">CONCEPT CHECK · {"교사용 / 정답 포함" if include_answers else "학생용 / 문제만"}</p><h1>{e(paper["title"])}</h1>')
        out.append(f'<p class="meta">{e(paper.get("subject",""))} · {e(paper.get("unit",""))} · {len(paper["questions"])}문항</p>')
        out.append('<p>학급: __________　이름: ______________　날짜: ______________</p>')
        for q in paper["questions"]:
            out.append(f'<article class="question"><div class="number">문항 {q.get("number",1):02d}</div><p class="prompt">{e(q["prompt"])}</p><ol class="options">')
            for j,opt in enumerate(q["options"],1):
                out.append(f'<li>{j}. {e(opt)}</li>')
            out.append('</ol><div class="workspace">나의 생각</div></article>')
        if include_answers:
            out.append('<section class="answer-section"><h2>정답 및 해설 · 교사용</h2>')
            for q in paper["questions"]:
                out.append(f'<article class="answer"><h3>{q.get("number",1):02d}번 · 정답 {q["answer"]+1}번</h3>'
                           f'<p>{e(q["options"][q["answer"]])}</p><p>{e(q.get("explanation",""))}</p></article>')
            out.append('</section>')
        out.append('</section>')
    out.append('</main><script>document.getElementById("print-quiz").addEventListener("click",function(){window.print()});</script></body></html>')
    return ''.join(out).encode('utf-8')

def build_export(papers, *, include_answers=False, format="docx"):
    _check(papers,include_answers)
    count=sum(len(p["questions"]) for p in papers)
    version="교사용_정답포함" if include_answers else "학생용_문제만"
    if len(papers)==1:
        name=filename(papers[0]["title"])
        if len(papers[0]["questions"])==1:
            name+=f'_문항{papers[0]["questions"][0].get("number",1):02d}'
    else: name=f"전체출제문항_{len(papers)}퀴즈_{count}문항"
    name=f"{name}_{version}"
    if format=="docx":
        data,mime,ext=_docx(papers,include_answers),DOCX_MIME,"docx"
    elif format=="html":
        data,mime,ext=_html(papers,include_answers),"text/html","html"
    elif format=="zip_docx":
        out=BytesIO()
        with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
            z.writestr("안내.txt",("교사가 배포했거나 종료한 퀴즈의 출제 문항입니다.\n"
                        f"전체 {len(papers)}퀴즈 · {count}문항 · {version}\n"
                        "각 퀴즈 폴더의 전체문항 파일과 개별문항 폴더를 사용하세요.\n"
                        "과거 퀴즈의 중복 유형은 당시 기록 보존을 위해 그대로 내보냅니다.\n"
                        "학생 이름·답안·성적·계정 정보는 포함하지 않습니다.\n").encode("utf-8"))
            for i,paper in enumerate(papers,1):
                folder=f'{i:03d}_{filename(paper["title"],48)}_{filename(paper["id"][:8])}'
                z.writestr(folder+f"/전체문항_{version}.docx",_docx([paper],include_answers))
                for q in paper["questions"]:
                    one=dict(paper,questions=[deepcopy(q)])
                    z.writestr(folder+f'/개별문항/문항{q.get("number",1):02d}_{version}.docx',_docx([one],include_answers))
        data,mime,ext=out.getvalue(),"application/zip","zip"
    else:
        raise ValueError("지원되는 파일 형식은 docx, html, zip_docx입니다.")
    return {"data":data,"filename":f"{name}.{ext}","mime":mime,"question_count":count,"quiz_count":len(papers)}
