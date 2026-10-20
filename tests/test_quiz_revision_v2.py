"""Variable-length quizzes, immutable legacy attempts and authorized downloads."""
from copy import deepcopy
from io import BytesIO
import csv
import json
import zipfile
from pathlib import Path
import pytest
from test_concept_quiz import env
from quiz_core import bank
from quiz_core.exports import build_export
from quiz_core.student_view import hero_html, question_html, given_visual_html

@pytest.mark.parametrize("subject",bank.SUBJECTS)
@pytest.mark.parametrize("index",range(8))
def test_declared_type_count_matches_generation(subject,index):
    unit=bank.units(subject)[index]
    assert 1<=bank.available_type_count(subject,unit)<=10
    for seed in range(5):
        qs=bank.generate(subject,unit,seed)
        assert len(qs)==bank.available_type_count(subject,unit)

@pytest.mark.parametrize("maximum",range(1,11))
def test_requested_maximum_is_not_padded(maximum):
    unit=bank.units("공통수학2")[5]
    qs=bank.generate("공통수학2",unit,42,max_questions=maximum)
    assert len(qs)==min(maximum,bank.available_type_count("공통수학2",unit))

@pytest.mark.parametrize("maximum",[0,11,-1,True,None,"5",3.0])
def test_invalid_maximum_rejected(maximum):
    with pytest.raises(ValueError):
        bank.generate("공통수학2",bank.units("공통수학2")[5],max_questions=maximum)

def test_number_changes_dont_make_new_type():
    qs=bank.generate("공통수학2",bank.units("공통수학2")[5],seed=12)
    other=deepcopy(qs[0]);other["prompt"] += " (수치를 바꾸어 출제)"
    with pytest.raises(ValueError,match="유형"):
        bank.validate_questions([qs[0],other])
    other["type_id"]="forged-new-type"
    with pytest.raises(ValueError,match="유형"):
        bank.validate_questions([qs[0],other])

def test_empty_and_overlimit_paper_rejected():
    with pytest.raises(ValueError):bank.validate_questions([])
    q=bank.generate("공통수학2",bank.units("공통수학2")[5])[0]
    with pytest.raises(ValueError):bank.validate_questions([q]*11)

def variable_paper(env,n):
    env.actor(env.t)
    qid=env.svc.create_draft(env.t,"공통수학2",bank.units("공통수학2")[5],max_questions=n)
    env.svc.publish(env.t,qid,["1-1"],reviewed=True)
    return qid

@pytest.mark.parametrize("n,correct,score",[(1,1,100),(3,1,33.3),(3,2,66.7),(6,5,83.3),(7,6,85.7),(7,0,0)])
def test_normalized_scores_and_dynamic_csv(env,n,correct,score):
    qid=variable_paper(env,n);u=env.students[0]
    paper=env.svc.teacher_paper(env.t,qid);questions=paper["questions"]
    choices=[q["answer"] if i<correct else (q["answer"]+1)%4 for i,q in enumerate(questions)]
    env.answer(qid,u,choices)
    env.svc.submit(u,qid)
    env.actor(env.t);env.svc.grade(env.t,qid,[u["id"]])
    board=env.svc.dashboard(env.t,qid)
    assert board["students"][0]["score"]==score
    data=list(csv.reader(BytesIO(env.svc.export_dashboard(env.t,qid)).getvalue().decode("utf-8-sig").splitlines()))
    assert len(data[0])==6+n
    assert len(data[1])==6+n
    env.actor(u)
    assert "score" not in env.svc.student_paper(u,qid)
    env.actor(env.t);env.svc.release(env.t,qid,[u["id"]]);env.actor(u)
    assert env.svc.student_paper(u,qid)["score"]==score
    assert env.svc.history(u,u["id"])[0]["score"]==score

def test_student_history_has_no_download_and_teacher_keeps_own(env):
    qid=env.paper();u=env.students[0];env.answer(qid,u)
    env.st.outputs.clear();env.view._history(u,u["id"],"student")
    assert not [o for o in env.st.outputs if o[0]=="download"]
    assert [o for o in env.st.outputs if o[0]=="dataframe"]
    env.actor(env.t);env.st.outputs.clear();env.view._history(env.t,u["id"],"teacher")
    assert any(o[0]=="download" and o[2]["mime"]=="application/json" for o in env.st.outputs)

def legacy_paper(env,publish=True):
    qid=env.svc.create_draft(env.t,"공통수학2",bank.units("공통수학2")[5])
    row=next(q for q in env.backend.read("quiz_sets") if q["id"]==qid)
    q=json.loads(row["question_json"])[0];old=[]
    for i in range(10):
        item=deepcopy(q);item.update(id=f"q{i+1:02}",prompt=f'{q["prompt"]} [기존 문항 {i+1}]')
        item.pop("type_id",None);old.append(item)
    env.backend.update("quiz_sets",qid,{"question_json":json.dumps(old,ensure_ascii=False),"generator":"concept-templates-1"})
    if publish:
        targets=[{k:env.students[0][k] for k in ("id","name","class")}]
        env.backend.update("quiz_sets",qid,{"state":"open","targets_json":json.dumps(targets),"classes_json":'["1-1"]'})
    return qid

def test_legacy_published_ten_remains_unchanged(env):
    qid=legacy_paper(env);old=deepcopy(env.svc.teacher_paper(env.t,qid)["questions"])
    u=env.students[0];env.answer(qid,u);env.svc.submit(u,qid)
    env.actor(env.t);env.svc.grade(env.t,qid,[u["id"]])
    assert env.svc.dashboard(env.t,qid)["students"][0]["score"]==100
    assert env.svc.teacher_paper(env.t,qid)["questions"]==old
    result=env.svc.export_questions(env.t,[qid],format="html")
    assert result["question_count"]==10
    assert result["data"].count(b'class="question"')==10
    with pytest.raises(ValueError):env.svc.deduplicate_draft(env.t,qid,env.svc.teacher_paper(env.t,qid)["revision"])

def test_legacy_draft_dedup_is_explicit(env):
    qid=legacy_paper(env,False)
    before=env.svc.teacher_paper(env.t,qid)
    assert len(before["questions"])==10
    n=env.svc.deduplicate_draft(env.t,qid,before["revision"])
    assert n==1
    env.svc.publish(env.t,qid,["1-1"],reviewed=True)

def test_download_permissions_even_forged_role(env):
    qid=env.paper()
    for actor in [env.students[0],env.other]:
        env.actor(actor)
        with pytest.raises(PermissionError):env.svc.export_questions(actor,[qid],format="html",include_answers=True)
    env.actor(env.students[0])
    with pytest.raises(PermissionError):env.svc.export_questions(dict(env.students[0],role="teacher"),[qid])
    with pytest.raises(PermissionError):env.view._question_download_view(env.students[0])

def test_multiple_owner_download_is_all_or_nothing(env):
    qid=env.paper();env.actor(env.other)
    other=env.svc.create_draft(env.other,"공통수학2",bank.units("공통수학2")[5])
    env.svc.publish(env.other,other,["1-1"],reviewed=True)
    env.actor(env.t)
    with pytest.raises(PermissionError):env.svc.export_questions(env.t,[qid,other],format="html")

def test_docx_single_whole_and_zip_no_student_data(env):
    from docx import Document
    qid=variable_paper(env,3);paper=env.svc.teacher_paper(env.t,qid)
    selected=paper["questions"][1]
    one=env.svc.export_questions(env.t,[qid],question_id=selected["id"])
    doc=Document(BytesIO(one["data"]))
    text="\n".join(p.text for p in doc.paragraphs)
    assert selected["prompt"] in text and paper["questions"][0]["prompt"] not in text
    assert "정답 및 해설" not in text
    assert one["question_count"]==1
    whole=env.svc.export_questions(env.t,[qid],include_answers=True)
    text="\n".join(p.text for p in Document(BytesIO(whole["data"])).paragraphs)
    assert "정답 및 해설" in text and all(q["explanation"] in text for q in paper["questions"])
    archive=env.svc.export_questions(env.t,[qid],format="zip_docx")
    with zipfile.ZipFile(BytesIO(archive["data"])) as z:
        assert len([x for x in z.namelist() if x.endswith(".docx")])==4
        assert all(".." not in Path(x).parts and not Path(x).is_absolute() for x in z.namelist())
        assert not any(x.endswith((".json",".csv")) for x in z.namelist())

def test_exports_include_closed_quizzes_and_no_student_pii(env):
    a=variable_paper(env,3);b=variable_paper(env,2)
    env.answer(a,env.students[0]);env.actor(env.t);env.svc.close(env.t,a)
    data=env.svc.export_questions(env.t,[a,b],format="html")["data"].decode()
    assert data.count('class="question"')==5
    assert env.students[0]["id"] not in data and env.students[0]["name"] not in data

def test_download_ui_preparation_is_explicit(env):
    qid=variable_paper(env,3)
    env.st.events.add("qexport_all")
    env.view._question_download_view(env.t)
    downloads=[o for o in env.st.outputs if o[0]=="download"]
    assert len(downloads)==1 and downloads[0][2]["file_name"].endswith(".docx")
    assert not env.st.session_state["qexport_answers"]
    env.st.session_state["qexport_answers"]=True
    env.st.outputs.clear()
    env.view._question_download_view(env.t)
    assert not [o for o in env.st.outputs if o[0]=="download"]

def test_html_exports_escape_script_and_no_keys_default(env):
    qid=variable_paper(env,2)
    row=env.backend.read("quiz_sets")[0]
    qs=json.loads(row["question_json"]);qs[0]["prompt"]='<script>window.BAD=1</script>'
    qs[0]["explanation"]="SECRET_EXPLANATION"
    env.backend.update("quiz_sets",qid,{"question_json":json.dumps(qs)})
    html=env.svc.export_questions(env.t,[qid],format="html")["data"].decode()
    assert "<script>window.BAD" not in html and "&lt;script&gt;" in html
    assert "SECRET_EXPLANATION" not in html
    assert "SECRET_EXPLANATION" in env.svc.export_questions(env.t,[qid],format="html",include_answers=True)["data"].decode()

def test_student_visuals_use_givens_not_answers():
    html=given_visual_html("명제 ‘x>7이면 x>3이다’의 대우는?")
    assert "x&gt;7" in html and "x&gt;3이다" in html
    assert "≤" not in html
    assert given_visual_html("원 (x-2)²+(y-3)²=25의 중심의 x좌표는?")==""
    html=given_visual_html("A(1, 2), B(5, 8)의 중점의 x좌표는?")
    assert "A(1, 2)" in html and "B(5, 8)" in html and "(3, 5)" not in html
    assert "qz-matrix" in given_visual_html("행렬 [[2, 3], [4, 5]]의 (2, 1) 성분은?")

def test_visuals_escape_untrusted_prompts():
    q={"id":"q01","prompt":'<img src=x onerror="BAD=1">',"options":["a","b","c","d"]}
    assert "<img" not in question_html(q,1,1,None)
    assert "&lt;img" in question_html(q,1,1,None)
    h=hero_html({"title":"<script>x</script>","subject":"수학","questions":[q],"answers":{"q01":0}})
    assert "<script>" not in h and "선택 1 / 1" in h
